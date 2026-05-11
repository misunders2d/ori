"""generate_presentation tests — layout dispatch, template fallback,
output path shape, validation error surfaces.

We stay within python-pptx's actual API: each test produces a real
``.pptx`` file in ``tmp_path`` and re-opens it to inspect slide count
and basic structure. Avoids mocking python-pptx (too easy to drift
from reality).
"""

from __future__ import annotations

import os
import pathlib

import pytest

from app.tools.presentations import (
    _resolve_template,
    _slugify,
    generate_presentation,
)


@pytest.fixture
def redirected_io(tmp_path, monkeypatch):
    """Redirect template + export directories to a per-test tmp dir so
    no cross-test leakage and no writes to the real data/."""
    templates = tmp_path / "templates"
    exports = tmp_path / "exports"
    templates.mkdir()
    exports.mkdir()
    monkeypatch.setattr("app.tools.presentations._TEMPLATES_DIR", str(templates))
    monkeypatch.setattr("app.tools.presentations._EXPORTS_DIR", str(exports))
    return {"templates": templates, "exports": exports}


# ---------------------------------------------------------------------------
# Slugify + output path helpers
# ---------------------------------------------------------------------------


def test_slugify_strips_special_chars():
    assert _slugify("Weekly Sales — Q2'26!") == "weekly-sales-q2-26"
    assert _slugify("   spaces   ") == "spaces"
    assert _slugify("") == "deck"


# ---------------------------------------------------------------------------
# Template resolution
# ---------------------------------------------------------------------------


def test_resolve_template_returns_none_when_no_files(redirected_io):
    assert _resolve_template("") is None
    assert _resolve_template("ghost") is None


def test_resolve_template_picks_default_implicitly(redirected_io):
    """An implicit lookup (template_name='') falls back to
    ``default.pptx``. This is the brand-template happy path: drop
    ``default.pptx`` and every deck inherits it."""
    default = redirected_io["templates"] / "default.pptx"
    default.write_bytes(b"x")  # contents don't matter for resolution
    assert _resolve_template("") == str(default)


def test_resolve_template_honours_explicit_name(redirected_io):
    """When ``template_name`` is passed, it wins over default."""
    (redirected_io["templates"] / "default.pptx").write_bytes(b"x")
    (redirected_io["templates"] / "executive.pptx").write_bytes(b"x")
    assert _resolve_template("executive") == str(
        redirected_io["templates"] / "executive.pptx"
    )


def test_resolve_template_falls_through_to_default_when_named_missing(redirected_io):
    """Named template not found → falls back to ``default.pptx`` rather
    than failing immediately. The tool still works; only the user's
    specific request is unfulfilled, and we log which template was
    actually used."""
    (redirected_io["templates"] / "default.pptx").write_bytes(b"x")
    assert _resolve_template("ghost") == str(
        redirected_io["templates"] / "default.pptx"
    )


# ---------------------------------------------------------------------------
# Validation — bad specs fail fast
# ---------------------------------------------------------------------------


def test_generate_presentation_rejects_empty_slides(redirected_io):
    res = generate_presentation(title="Deck", slides=[], tool_context=None)
    assert res["status"] == "error"
    assert "non-empty" in res["message"]


def test_generate_presentation_rejects_unknown_layout(redirected_io):
    res = generate_presentation(
        title="Deck",
        slides=[{"layout": "barchart", "title": "x"}],
        tool_context=None,
    )
    assert res["status"] == "error"
    assert "unknown layout" in res["message"]
    assert "barchart" in res["message"]


def test_generate_presentation_rejects_non_dict_slide(redirected_io):
    res = generate_presentation(
        title="Deck",
        slides=[{"layout": "title", "title": "x"}, "not a dict"],
        tool_context=None,
    )
    assert res["status"] == "error"
    assert "slide 1" in res["message"]


# ---------------------------------------------------------------------------
# Successful generation — verify file shape + slide count
# ---------------------------------------------------------------------------


def _open_deck(path: str):
    from pptx import Presentation

    return Presentation(path)


def test_generate_minimal_title_deck(redirected_io):
    res = generate_presentation(
        title="Hello",
        slides=[{"layout": "title", "title": "Hello", "subtitle": "world"}],
        tool_context=None,
    )
    assert res["status"] == "success"
    assert res["slide_count"] == 1
    assert res["template_used"] == "blank"  # no template at all in tmp dir
    assert os.path.isfile(res["file_path"])
    assert res["file_path"].endswith(".pptx")

    deck = _open_deck(res["file_path"])
    assert len(deck.slides) == 1


def test_generate_multi_layout_deck(redirected_io, tmp_path):
    """One slide per layout type — sanity check that the dispatch table
    + helper functions all execute without raising on a representative
    spec for each."""
    # Make a tiny PNG so chart / image slides have something to embed.
    img = tmp_path / "chart.png"
    img.write_bytes(
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc\xfc"
        b"\xcf\xc0\xf0\x1f\x00\x05\x00\x01\xff\xa5\xd5\x9e\xae\x00\x00\x00"
        b"\x00IEND\xaeB`\x82"
    )

    slides = [
        {"layout": "title", "title": "Deck", "subtitle": "subtitle"},
        {"layout": "bullets", "title": "List", "bullets": ["a", "b", "c"]},
        {"layout": "chart", "title": "Chart", "chart_path": str(img), "caption": "cap"},
        {
            "layout": "kpi_grid",
            "title": "Metrics",
            "kpis": [
                {"label": "Revenue", "value": "$100", "delta": "+5%"},
                {"label": "Units", "value": "200", "delta": "-2%"},
                {"label": "Conv", "value": "3.2%"},
            ],
        },
        {
            "layout": "table",
            "title": "Top ASINs",
            "columns": ["ASIN", "Units"],
            "rows": [["B001", "100"], ["B002", "80"]],
        },
        {"layout": "two_column", "title": "vs", "left": "before", "right": "after"},
        {"layout": "image", "title": "Photo", "image_path": str(img)},
        {"layout": "text", "title": "Summary", "paragraph": "Long-form text body."},
    ]
    res = generate_presentation(title="Full", slides=slides, tool_context=None)
    assert res["status"] == "success", res
    assert res["slide_count"] == 8

    deck = _open_deck(res["file_path"])
    assert len(deck.slides) == 8


def test_generate_deck_with_template(redirected_io):
    """When ``default.pptx`` exists under the templates dir, the tool
    must use it as the starting Presentation and report the path in
    ``template_used``. We seed the templates dir with a real (blank)
    pptx so python-pptx can actually open it."""
    from pptx import Presentation

    template_path = redirected_io["templates"] / "default.pptx"
    Presentation().save(str(template_path))

    res = generate_presentation(
        title="Branded",
        slides=[{"layout": "title", "title": "Branded"}],
        tool_context=None,
    )
    assert res["status"] == "success"
    assert res["template_used"] == str(template_path)


def test_generate_deck_handles_missing_chart_path_gracefully(redirected_io):
    """A chart slide whose ``chart_path`` is empty / nonexistent should
    still render (title + caption) rather than abort the whole deck."""
    res = generate_presentation(
        title="Missing chart",
        slides=[
            {"layout": "title", "title": "OK"},
            {
                "layout": "chart",
                "title": "where's the chart",
                "chart_path": "/tmp/nonexistent_chart_42.png",
                "caption": "still here",
            },
        ],
        tool_context=None,
    )
    assert res["status"] == "success"
    assert res["slide_count"] == 2


def test_filename_kwarg_used_verbatim(redirected_io):
    res = generate_presentation(
        title="x",
        slides=[{"layout": "title", "title": "x"}],
        filename="my-report",
        tool_context=None,
    )
    assert res["file_path"].endswith("my-report.pptx")


def test_filename_kwarg_keeps_pptx_when_already_present(redirected_io):
    res = generate_presentation(
        title="x",
        slides=[{"layout": "title", "title": "x"}],
        filename="my-report.pptx",
        tool_context=None,
    )
    assert res["file_path"].endswith("my-report.pptx")
    # No double-extension
    assert not res["file_path"].endswith(".pptx.pptx")


# ---------------------------------------------------------------------------
# Broken template handling
# ---------------------------------------------------------------------------


def test_generate_deck_fails_loudly_on_corrupt_template(redirected_io):
    """A corrupt template file is an explicit error, not a silent
    fallback to blank — the user needs to know their brand template is
    broken instead of getting an unbranded deck and wondering why."""
    bad = redirected_io["templates"] / "default.pptx"
    bad.write_text("this is not a valid pptx file")

    res = generate_presentation(
        title="x",
        slides=[{"layout": "title", "title": "x"}],
        tool_context=None,
    )
    assert res["status"] == "error"
    assert "template" in res["message"].lower()
